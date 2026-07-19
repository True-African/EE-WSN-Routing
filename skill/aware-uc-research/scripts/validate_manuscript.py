import argparse
from pathlib import Path

from docx import Document


REPO_ROOT = Path(__file__).resolve().parents[3]
MANUSCRIPT_DIR = REPO_ROOT / "manuscript"


def latest_docx():
    candidates = sorted(
        MANUSCRIPT_DIR.glob("AWARE-UC*.docx"),
        key=lambda p: p.stat().st_mtime,
        reverse=True,
    )
    if not candidates:
        raise FileNotFoundError("No generated AWARE-UC manuscript was found in manuscript/.")
    return candidates[0]


def main():
    parser = argparse.ArgumentParser(description="Validate the generated AWARE-UC Word manuscript.")
    parser.add_argument("path", nargs="?", type=Path, help="Optional manuscript .docx path")
    args = parser.parse_args()
    path = args.path.resolve() if args.path else latest_docx()
    doc = Document(path)
    text = "\n".join(paragraph.text for paragraph in doc.paragraphs)

    checks = {
        "AWARE-UC name": "AWARE-UC" in text,
        "Original flow chart discussed": "Original thesis routing/activity diagram" in text,
        "New flow chart caption": "Figure 6. Proposed AWARE-UC" in text,
        "Pseudocode present": "Algorithm 1: AWARE-UC Pseudocode" in text,
        "Radio equation present": "E_TX(k,d)" in text,
        "Threshold equation present": "E_th(r)" in text,
        "CH score equation present": "Score_i(r)" in text,
        "SEP baseline present": "SEP" in text,
        "Recent comparison table present": "Recent Protocol Comparison" in text,
        "Symbol table present": "Symbol Table" in text,
        "Complexity analysis present": "Computational Complexity" in text,
    }

    doi_refs = sum(1 for paragraph in doc.paragraphs if "doi.org" in paragraph.text.lower())

    print(f"Validated file: {path}")
    print(f"Paragraphs: {len(doc.paragraphs)}")
    print(f"Tables: {len(doc.tables)}")
    print(f"Figures/images: {len(doc.inline_shapes)}")
    print(f"DOI reference paragraphs: {doi_refs}")
    print()

    failed = []
    for name, ok in checks.items():
        status = "PASS" if ok else "FAIL"
        print(f"{status}: {name}")
        if not ok:
            failed.append(name)

    print()
    if doi_refs < 30:
        failed.append("At least 30 DOI-bearing reference paragraphs")
        print("FAIL: At least 30 DOI-bearing reference paragraphs")
    else:
        print("PASS: At least 30 DOI-bearing reference paragraphs")

    if len(doc.inline_shapes) < 12:
        failed.append("Expected figure/image count")
        print("FAIL: Expected figure/image count")
    else:
        print("PASS: Expected figure/image count")

    if failed:
        raise SystemExit(f"Validation failed: {', '.join(failed)}")

    print("\nValidation passed.")


if __name__ == "__main__":
    main()
